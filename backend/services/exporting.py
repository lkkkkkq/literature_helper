import base64
import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from backend.core.errors import DomainError


def safe_cell(value):
    if not isinstance(value, str):
        return "" if value is None else str(value)
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r")):
        return "'" + text
    return text


def excel_cell(value):
    if isinstance(value, str):
        if len(value) > 32766:
            raise DomainError("单个字段超过 Excel 单元格限制，请使用 CSV 导出。")
        return safe_cell(value)
    return value


def export_snapshot(snapshot, format):
    schema, job = snapshot["schema"], snapshot["job"]
    fields = schema["plan"]["fields"]
    headers = [
        "record_id",
        "document_id",
        "job_id",
        "schema_version",
        "review_version",
        "risks",
        "metadata",
    ]
    for field in fields:
        headers.extend(
            [
                field["key"],
                field["key"] + "__unit",
                field["key"] + "__missing",
                field["key"] + "__evidence",
                field["key"] + "__standardized",
            ]
        )
    rows = []
    evidence_rows = [["record_id", "field", "file", "sha256", "page", "bbox", "type", "quote"]]
    for record in snapshot["records"]:
        row = [
            record["id"],
            record["document_id"],
            job["id"],
            schema["version"],
            snapshot["version"],
            json.dumps(record["risks"], ensure_ascii=False),
            json.dumps(
                {
                    "project_id": job["project_id"],
                    "job_status": job["status"],
                    "schema": schema,
                    "history": [
                        item
                        for item in snapshot["history"][: snapshot["version"]]
                        if item["record_id"] == record["id"]
                    ],
                },
                ensure_ascii=False,
            )
            if format == "csv"
            else "See Metadata sheet",
        ]
        for field in fields:
            value = record["values_json"][field["key"]]
            evidence = [
                item for item in record["evidence_json"] if item["id"] in value["evidence_ids"]
            ]
            row.extend(
                [
                    value["value"],
                    value["unit"],
                    value["missing_reason"],
                    json.dumps(
                        evidence if format == "csv" else value["evidence_ids"], ensure_ascii=False
                    ),
                    json.dumps(value.get("standardized"), ensure_ascii=False),
                ]
            )
            for item in evidence:
                for location in item["locations"] or [{"page": "", "bbox": None}]:
                    evidence_rows.append(
                        [
                            record["id"],
                            field["key"],
                            item["filename"],
                            item["source_sha256"],
                            location["page"],
                            json.dumps(location["bbox"]),
                            item["kind"],
                            item["quote"],
                        ]
                    )
        rows.append(row)
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows([[safe_cell(cell) for cell in row] for row in rows])
        data = output.getvalue().encode("utf-8-sig")
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Records"
        for row in [headers, *rows]:
            sheet.append([excel_cell(cell) for cell in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        evidence_sheet = book.create_sheet("Evidence")
        for row in evidence_rows:
            evidence_sheet.append([excel_cell(cell) for cell in row])
        evidence_sheet.freeze_panes = "A2"
        metadata = book.create_sheet("Metadata")
        for key, value in {
            "project_id": job["project_id"],
            "job_id": job["id"],
            "job_status": job["status"],
            "schema": schema,
            "review_version": snapshot["version"],
            "history": snapshot["history"][: snapshot["version"]],
        }.items():
            serialized = json.dumps(value, ensure_ascii=False)
            for index in range(0, len(serialized), 30000):
                metadata.append(
                    [key, index // 30000 + 1, safe_cell(serialized[index : index + 30000])]
                )
        for worksheet in book:
            for cells in worksheet:
                for cell in cells:
                    if isinstance(cell.value, str) and len(cell.value) > 32767:
                        raise DomainError("单个字段超过 Excel 单元格限制，请使用 CSV 导出。")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24557A")
            for column in worksheet.columns:
                worksheet.column_dimensions[column[0].column_letter].width = 28
        evidence_sheet.column_dimensions["H"].width = 80
        metadata.column_dimensions["C"].width = 100
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
    return {
        "filename": f"litextract-{job['id'][:8]}-v{snapshot['version']}.{format}",
        "data": base64.b64encode(data).decode("ascii"),
    }
