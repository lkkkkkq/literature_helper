import base64
import csv
import io
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from test_chat import chat  # noqa: F401
from test_jobs import plan
from test_reading import reading  # noqa: F401


@pytest.mark.parametrize("with_document_id", [False, True])
def test_reviews_are_append_only_scoped_and_export_exact_versions(chat, with_document_id):  # noqa: F811
    client, answer_path, model, project_path, _ = chat
    service = client.app.state.jobs
    service.close()
    schema = plan()
    if with_document_id:
        schema["fields"].append(
            {"key": "source_document_id", "label": "源文档 ID", "type": "string"}
        )
    model.complete = lambda *args: schema
    base = answer_path.removesuffix("/answers")
    draft = client.post(
        base + "/plans", json={"goal": "提取站点"}, headers={"X-Model-Key": "test"}
    ).json()
    job = client.post(
        base + "/jobs",
        json={"schema_id": draft["id"], "request_id": str(uuid4()), "plan": schema},
        headers={"X-Model-Key": "test"},
    ).json()
    path = project_path + f"/jobs/{job['id']}"
    detail = client.get(path).json()
    pid = project_path.split("/")[-1]
    did = detail["documents"][0]["document_id"]
    source = client.app.state.chat.sources(pid, [did])[0][0].model_dump()
    repo = service.repository
    repo.claim(pid, job["id"])
    repo.begin_document(pid, job["id"], did)
    repo.finish_document(
        pid,
        job["id"],
        did,
        [
            {
                "values_json": {
                    **(
                        {
                            "source_document_id": {
                                "value": did,
                                "unit": "",
                                "missing_reason": "",
                                "evidence_ids": [],
                            }
                        }
                        if with_document_id
                        else {}
                    ),
                    "site": {
                        "value": "Alpha",
                        "unit": "",
                        "missing_reason": "",
                        "evidence_ids": [source["id"]],
                    },
                },
                "evidence_json": [source],
                "risks": [],
            }
        ],
        [],
    )
    repo.finish(pid, job["id"])
    original = client.get(path + "/results").json()
    row = original["records"][0]
    if with_document_id:
        response = client.post(
            path + "/reviews",
            json={
                "base_version": 0,
                "record_id": row["id"],
                "field_key": "source_document_id",
                "after": {"value": "invented"},
                "reason": "尝试篡改来源",
            },
        )
        assert response.status_code == 400
        assert "系统文献索引" in response.text
    body = {
        "base_version": 0,
        "record_id": row["id"],
        "field_key": "site",
        "after": {**row["values_json"]["site"], "value": "=SUM(1,2)"},
        "reason": "人工核对",
    }
    assert client.post(path + "/reviews", json=body).status_code == 201
    assert client.post(path + "/reviews", json=body).status_code == 409
    assert (
        client.get(path + "/results?version=automatic").json()["records"][0]["values_json"]["site"][
            "value"
        ]
        == "Alpha"
    )
    assert (
        client.get(path + "/results?version=1").json()["records"][0]["values_json"]["site"]["value"]
        == "=SUM(1,2)"
    )
    assert client.get(path + "/results?version=2").status_code == 404
    assert client.get(path.replace(pid, str(uuid4())) + "/results").status_code == 404
    body["base_version"] = 1
    body["after"]["evidence_ids"] = ["invented"]
    assert client.post(path + "/reviews", json=body).status_code == 400
    for format in ["csv", "xlsx"]:
        result = client.get(path + f"/export?format={format}&version=1").json()
        data = base64.b64decode(result["data"])
        if format == "csv":
            exported = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
            assert exported[0]["site"] == "'=SUM(1,2)"
            assert source["quote"] in exported[0]["site__evidence"]
            assert exported[0]["review_version"] == "1"
        else:
            book = load_workbook(io.BytesIO(data))
            assert book.sheetnames == ["Records", "Evidence", "Metadata"]
            assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
            assert book["Evidence"].max_row > 1
    assert client.get(path + "/records").json()[0]["values_json"]["site"]["value"] == "Alpha"
