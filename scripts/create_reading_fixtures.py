"""Generate original, deterministic documents for reading acceptance tests."""

import argparse
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle


def create_fixtures(output: Path):
    output.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(output / "paper.pdf"), pagesize=(612, 792))
    pdf.setTitle("River study - reading fixture")
    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawString(54, 720, "River monitoring study")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(54, 680, "Methods: samples were collected from two river sites.")
    pdf.showPage()
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(54, 720, "Results")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(54, 680, "The observed concentration was 12.5 mg/L at site Alpha.")
    table = Table(
        [["Site", "Value (mg/L)"], ["Alpha", "12.5"], ["Beta", "8.2"]],
        colWidths=[170, 170],
        rowHeights=28,
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    table.wrapOn(pdf, 500, 500)
    table.drawOn(pdf, 54, 540)
    pdf.save()
    (output / "supp.csv").write_text("Site,Value\nAlpha,12.5\nBeta,8.2\n", encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    for row in [["Site", "Value"], ["Alpha", 12.5], ["Beta", 8.2]]:
        sheet.append(row)
    workbook.save(output / "supp.xlsx")
    document = Document()
    document.add_heading("Supplementary methods", level=1)
    document.add_paragraph("Samples were stored at 4 degrees Celsius.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Site", "Value"
    row = table.add_row().cells
    row[0].text, row[1].text = "Alpha", "12.5"
    document.save(output / "supp.docx")
    image = Image.new("RGB", (1224, 1584), "white")
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 38)
    draw = ImageDraw.Draw(image)
    draw.text((100, 140), "Scanned river sample", fill="black", font=font)
    draw.text((100, 240), "Alpha concentration 12.5 mg/L", fill="black", font=font)
    chinese = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 38)
    draw.text((100, 340), "河流样品浓度", fill="black", font=chinese)
    image.save(output / "scanned.pdf", "PDF", resolution=144)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    create_fixtures(parser.parse_args().output)
