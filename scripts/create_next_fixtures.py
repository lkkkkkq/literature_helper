"""Explicitly constructed fixtures, not real-paper accuracy evidence."""

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw


def create(root):
    root.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (400, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(40, 30), (40, 270), (360, 270)], fill="black", width=2)
    draw.line([(40, 270), (200, 150), (360, 30)], fill="red", width=3)
    image.save(root / "fixture-plot.png")
    document = Document()
    document.add_heading("Constructed native document fixture", 0)
    document.add_paragraph("Known abundance: 0.00004 particles/L. This is synthetic test data.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Unit"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "particles/L"
    table.cell(1, 1).text = "0.00004"
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    document.add_paragraph()._p.append(math)
    document.add_picture(str(root / "fixture-plot.png"))
    document.add_paragraph("Figure 1. Constructed line for manual calibration.")
    document.save(root / "native.docx")
    (root / "main.tex").write_text(
        r"\documentclass{article}\begin{document}\input{section}"
        r"\includegraphics{fixture-plot.png}\end{document}",
        encoding="utf-8",
    )
    (root / "section.tex").write_text(
        r"\section{Synthetic} $P=\frac{T^{3/2}}{\sqrt{2\rho A}}$", encoding="utf-8"
    )


if __name__ == "__main__":
    create(Path(sys.argv[1]))
